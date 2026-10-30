#!/usr/bin/env python3
"""Export a manuscript for Word review, or extract a returned review.

Run with Python containing python-docx (export only), and Pandoc on PATH.
This prepares evidence for an agent-assisted merge; it never rewrites LaTeX.
"""

from __future__ import annotations

import argparse
import difflib
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
import tomllib
import uuid
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS = {"w": W, "m": M}


def run(args, *, text=None, cwd=None):
    """Run a checked subprocess with UTF-8 input/output."""
    result = subprocess.run(
        args,
        input=text,
        text=True,
        encoding="utf-8",
        capture_output=True,
        check=False,
        cwd=cwd,
    )
    if result.returncode:
        raise RuntimeError(f"Command failed: {args!r}\n{result.stderr}")
    if result.stderr:
        print(result.stderr, file=sys.stderr, end="")
    return result.stdout


def digest(data):
    return hashlib.sha256(data).hexdigest()


def inlines(text):
    return [{"t": "Str", "c": text}]


def plain(node):
    if isinstance(node, list):
        return "".join(plain(x) for x in node)
    if not isinstance(node, dict):
        return ""
    t, c = node.get("t"), node.get("c", [])
    if t in ("Str", "Code"):
        return c if t == "Str" else c[1]
    if t in ("Space", "SoftBreak", "LineBreak"):
        return " "
    if t == "Math":
        return c[1]
    if t in ("Span", "Link", "Image", "Cite"):
        return plain(c[1])
    if t == "Header":
        return plain(c[2])
    return plain(c)


def strip_comments(text):
    """Remove TeX comments but retain escaped percent signs."""
    text = re.sub(r"^[ \t]*%[^\n]*\n", "", text, flags=re.M)
    return re.sub(r"(?<!\\)%[^\n]*", "", text)


def brace_argument(text, start):
    depth = 0
    for i in range(start, len(text)):
        if text[i] in "{}" and (i == 0 or text[i - 1] != "\\"):
            depth += 1 if text[i] == "{" else -1
            if depth == 0:
                return text[start + 1 : i], i + 1
    raise ValueError("Unbalanced LaTeX argument")


def normalize_tex(tex, labels=None):
    """Expose journal-only headings and declarations to Pandoc."""
    sources = []

    def source_marker(match):
        path = match.group(1).strip()
        if Path(path).is_absolute():
            try:
                path = str(Path(path).relative_to(ROOT))
            except ValueError:
                pass
        sources.append(path)
        return f"\n\nMMSOURCE{len(sources):05d}TOKEN\n\n"

    tex = re.sub(r"^[ \t]*% Source file: (.+)$", source_marker, tex, flags=re.M)
    tex = strip_comments(tex)
    # Pandoc's LaTeX reader can consume the following hyphen and species number
    # as an argument to the unknown nobreakdash command (e.g. CFC-12 -> CFC).
    tex = tex.replace(r"\nobreakdash", "").replace(r"\protect", "")
    for command, title in {
        "codedataavailability": "Code and data availability",
        "codeavailability": "Code availability",
        "dataavailability": "Data availability",
        "authorcontribution": "Author contributions",
        "competinginterests": "Competing interests",
        "disclaimer": "Disclaimer",
    }.items():
        pattern = re.compile(r"\\" + command + r"\s*\{")
        while match := pattern.search(tex):
            body, end = brace_argument(tex, match.end() - 1)
            tex = (
                tex[: match.start()]
                + "\n\\section*{"
                + title
                + "}\n"
                + body
                + tex[end:]
            )
    tex = re.sub(r"\\introduction\b", r"\\section{Introduction}", tex)
    tex = re.sub(r"\\conclusions\b", r"\\section{Conclusions}", tex)
    tex = tex.replace(r"\begin{acknowledgements}", r"\section*{Acknowledgements}")
    tex = tex.replace(r"\end{acknowledgements}", "")
    pattern = re.compile(r"\\belowtable\s*\{")
    while match := pattern.search(tex):
        body, end = brace_argument(tex, match.end() - 1)
        tex = tex[: match.start()] + "\n\nTable note: " + body + tex[end:]
    # Make abstract a body section: Pandoc's DOCX writer omits abstract metadata.
    tex = tex.replace(r"\begin{abstract}", r"\section*{Abstract}")
    tex = tex.replace(r"\end{abstract}", "")
    tex = tex.replace("{table*}", "{table}")

    # A source marker or size command inside a float can make Pandoc silently
    # discard its caption. Expose table captions as explicit body paragraphs.
    def table_caption(match):
        body = match.group(1)
        caption_match = re.search(r"\\caption\s*\{", body)
        if not caption_match:
            return body
        caption, end = brace_argument(body, caption_match.end() - 1)
        label_match = re.search(r"\\label\{([^}]+)\}", body[end:])
        label = label_match[1] if label_match else ""
        number = (labels or {}).get(label, "")
        body = body[: caption_match.start()] + body[end:]
        if label:
            body = body.replace(r"\label{" + label + "}", "", 1)
        heading = r"\hypertarget{" + label + "}{} " if label else ""
        return (
            "\n\n"
            + heading
            + r"\textbf{Table "
            + number
            + ".} "
            + caption
            + "\n\n"
            + body.replace(r"\footnotesize", "")
            + "\n\n"
        )

    tex = re.sub(
        r"\\begin\{table\}(?:\[[^]]*\])?(.*?)\\end\{table\}",
        table_caption,
        tex,
        flags=re.S,
    )
    tex = re.sub(r"\\(?:top|middle|bottom)hline\b", r"\\hline", tex)
    return tex, sources


def export(args):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    name = args.manuscript + "-ghg-forcing-for-cmip7"
    build = ROOT / "build" / name / "latex"
    main = build / "main.tex"
    if not main.exists() or not main.with_suffix(".aux").exists():
        raise SystemExit(
            "Compile the manuscript first (main.tex and main.aux are required)."
        )
    export_id = (
        datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        + "-"
        + uuid.uuid4().hex[:8]
    )
    bundle = ROOT / "review-baselines" / name / export_id
    bundle.mkdir(parents=True)
    output = (
        args.output or ROOT / "compiled-manuscripts" / f"{name}_MM.docx"
    ).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    metadata_file = ROOT / "manuscripts" / name / "metadata.toml"
    metadata = tomllib.loads(metadata_file.read_text())
    labels = dict(
        re.findall(
            r"\\newlabel\{([^}]+)\}\{\{([^}]+)\}", main.with_suffix(".aux").read_text()
        )
    )
    tex, sources = normalize_tex(main.read_text(), labels)
    ast = json.loads(run(["pandoc", "-f", "latex", "-t", "json"], text=tex, cwd=build))
    provenance = {}
    state = {"source": None, "section": "Title", "index": 0, "headings": [0] * 6}
    drop = object()

    def transform(node):
        if isinstance(node, list):
            result = []
            for item in node:
                item = transform(item)
                if item is not drop:
                    result.append(item)
            return result
        if not isinstance(node, dict):
            return node
        t, c = node.get("t"), node.get("c")
        if t == "Para" and (
            marker := re.fullmatch(r"MMSOURCE(\d+)TOKEN", plain(node).strip())
        ):
            state["source"] = sources[int(marker[1]) - 1]
            return drop
        if t == "Header" and "unnumbered" not in c[1][1]:
            level = c[0] - 1
            state["headings"][level] += 1
            state["headings"][level + 1 :] = [0] * (5 - level)
            number = labels.get(c[1][0]) or ".".join(
                str(n) for n in state["headings"][: level + 1]
            )
            c[2] = inlines(number + " ") + c[2]
        if t in ("Figure", "Table") and c[0][0] in labels and c[1][1]:
            caption = c[1][1][0]
            if caption["t"] in ("Para", "Plain"):
                kind = "Figure" if t == "Figure" else "Table"
                caption["c"] = inlines(f"{kind} {labels[c[0][0]]}. ") + caption["c"]
        if t == "Link":
            attributes = dict(c[0][2])
            if "reference" in attributes:
                key = attributes["reference"]
                if key not in labels:
                    raise ValueError(f"Reference missing from compiled .aux: {key}")
                c[1] = inlines(labels[key])
        if t == "Image":
            source = build / c[2][0]
            if source.suffix.lower() == ".pdf":
                target = bundle / "media" / (source.stem + ".png")
                target.parent.mkdir(exist_ok=True)
                run(
                    [
                        "pdftoppm",
                        "-singlefile",
                        "-r",
                        "200",
                        "-png",
                        str(source),
                        str(target.with_suffix("")),
                    ]
                )
                c[2][0] = str(target)
                c[0][2] = [(k, v) for k, v in c[0][2] if k not in ("width", "height")]
                c[0][2].append(["width", "6.3in"])
        if "c" in node:
            node["c"] = transform(c)
        if t in ("Para", "Plain", "Header"):
            if t in ("Para", "Plain"):
                equations = [
                    child
                    for child in node["c"]
                    if child.get("t") == "Math" and child["c"][0]["t"] == "DisplayMath"
                ]
                equation_numbers = [
                    labels[key]
                    for equation in equations
                    for key in re.findall(r"\\label\{([^}]+)\}", equation["c"][1])
                    if key in labels
                ]
                if equation_numbers:
                    node["c"].extend(inlines(" (" + ", ".join(equation_numbers) + ")"))
            text = plain(node).strip()
            if not text:
                return node
            if t == "Header":
                state["section"] = text
            state["index"] += 1
            bookmark = f"MMR{state['index']:06d}"
            provenance[bookmark] = {
                "source": state["source"],
                "section": state["section"],
                "ast_text": text,
            }
            children = node["c"][2] if t == "Header" else node["c"]
            wrapped = [{"t": "Span", "c": [[bookmark, [], []], children]}]
            if t == "Header":
                node["c"][2] = wrapped
            else:
                node["c"] = wrapped
        return node

    ast["blocks"] = transform(ast["blocks"])
    title = metadata["title_info"]["title"]
    ast["meta"] = {"title": {"t": "MetaInlines", "c": inlines(title)}}
    # Use the actual metadata to retain journal-specific Author/affil commands.
    author_latex = ", ".join(
        a["given_name"]
        + " "
        + a["surname"]
        + r"\textsuperscript{"
        + ",".join(
            str(list(metadata["affiliations"]).index(key) + 1)
            for key in a["affiliations"]
        )
        + "}"
        for a in metadata["authors"]
    )
    author_ast = json.loads(
        run(["pandoc", "-f", "latex", "-t", "json"], text=author_latex)
    )
    affiliation_text = "\n\n".join(
        f"{i}. {text}" for i, text in enumerate(metadata["affiliations"].values(), 1)
    )
    front = (
        author_ast["blocks"]
        + json.loads(
            run(["pandoc", "-f", "markdown", "-t", "json"], text=affiliation_text)
        )["blocks"]
    )
    for author in metadata["authors"]:
        if author.get("correspondence_author"):
            front.append(
                {
                    "t": "Para",
                    "c": inlines(
                        "Correspondence: "
                        + author["given_name"]
                        + " "
                        + author["surname"]
                        + " ("
                        + author["email"]
                        + ")"
                    ),
                }
            )
    ast["blocks"] = front + ast["blocks"]
    # Materialize citations; preserve original keys in the source snapshots.
    cited = run(
        [
            "pandoc",
            "-f",
            "json",
            "-t",
            "json",
            "--citeproc",
            "--bibliography",
            str(build / "references.bib"),
            "-M",
            "reference-section-title=References",
        ],
        text=json.dumps(ast),
        cwd=build,
    )
    (bundle / "export.json").write_text(cited)
    with tempfile.TemporaryDirectory() as tmp:
        reference = Path(tmp) / "reference.docx"
        run(["pandoc", "-f", "markdown", "-o", str(reference)], text="Reference")
        refdoc = Document(reference)
        for section in refdoc.sections:
            section.page_width, section.page_height = Inches(8.5), Inches(11)
            section.top_margin = section.bottom_margin = Inches(0.75)
            section.left_margin = section.right_margin = Inches(0.85)
        for style in refdoc.styles:
            if style.type == 1:
                style.font.name = "Times New Roman"
                style.font.color.rgb = RGBColor(0, 0, 0)
                style.paragraph_format.space_after = Pt(6)
        for s in ("Normal", "Body Text", "First Paragraph"):
            refdoc.styles[s].font.size = Pt(12)
            refdoc.styles[s].paragraph_format.line_spacing = 1.1
        refdoc.styles["Title"].font.size = Pt(18)
        refdoc.styles["Caption"].font.size = Pt(11)
        refdoc.save(reference)
        run(
            [
                "pandoc",
                "-f",
                "json",
                "-o",
                str(output),
                "--reference-doc",
                str(reference),
            ],
            text=cited,
            cwd=build,
        )
    doc = Document(output)
    tracking = OxmlElement("w:trackRevisions")
    doc.settings.element.append(tracking)
    doc.core_properties.identifier = export_id
    doc.core_properties.subject = "Word review of " + name
    commit = run(["git", "rev-parse", "HEAD"], cwd=ROOT).strip()
    branch = run(["git", "branch", "--show-current"], cwd=ROOT).strip()
    doc.core_properties.comments = (
        f"Source commit {commit}; branch {branch}; export {export_id}"
    )
    for shape in doc.inline_shapes:
        scale = min(1, Inches(6.3) / shape.width, Inches(6.1) / shape.height)
        shape.width, shape.height = int(shape.width * scale), int(shape.height * scale)
    for table in doc.tables:
        table.autofit = True
        table_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + edge)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                el.set(qn("w:" + key), value)
            borders.append(el)
        table_pr.append(borders)
        for ri, row in enumerate(table.rows):
            if ri == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run_item in paragraph.runs:
                        run_item.font.size = Pt(10)
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = 1
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        p._p.append(field)
    doc.save(output)
    # Include every final Word paragraph in the map, including bibliography/frontmatter.
    paragraphs = read_docx(output)["paragraphs"]
    in_references = False
    for para in paragraphs:
        in_references = in_references or para["text"].strip() == "References"
        para.update(
            provenance.get(
                para["bookmark"],
                {
                    "source": (
                        "references/references.bib"
                        if in_references
                        else str(metadata_file.relative_to(ROOT))
                    ),
                    "section": "Metadata or references",
                },
            )
        )
    snapshots = {}
    for path in [
        *sorted((ROOT / "manuscripts" / name).glob("*")),
        ROOT / "references" / "references.bib",
    ]:
        if path.is_file():
            relative = str(path.relative_to(ROOT))
            snapshots[relative] = {
                "sha256": digest(path.read_bytes()),
                "text": path.read_text(),
            }
    manifest = {
        "export_id": export_id,
        "manuscript": name,
        "source_commit": commit,
        "source_branch": branch,
        "output": str(output),
        "paragraphs": paragraphs,
        "snapshots": snapshots,
        "labels": labels,
    }
    (bundle / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2)
    )
    shutil.copy2(output, bundle / "baseline.docx")
    shutil.copy2(main, bundle / "main.tex")
    shutil.copy2(main.with_suffix(".aux"), bundle / "main.aux")
    print(
        json.dumps(
            {
                "output": str(output),
                "baseline": str(bundle),
                "paragraphs": len(paragraphs),
                "mapped_bookmarks": sum(
                    p["bookmark"] in provenance for p in paragraphs
                ),
            },
            indent=2,
        )
    )


def read_docx(path):
    """Read accepted text, original text, revisions and anchored comments."""

    def text_of(el, mode="accepted"):
        if el.tag in (f"{{{W}}}del", f"{{{W}}}moveFrom") and mode == "accepted":
            return ""
        if el.tag in (f"{{{W}}}ins", f"{{{W}}}moveTo") and mode == "original":
            return ""
        if el.tag in (f"{{{W}}}t", f"{{{W}}}delText", f"{{{M}}}t"):
            return el.text or ""
        if el.tag in (f"{{{W}}}tab", f"{{{W}}}br"):
            return " "
        return "".join(text_of(c, mode) for c in el)

    with ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
        core = ET.fromstring(z.read("docProps/core.xml"))
        identifier = core.find("{http://purl.org/dc/elements/1.1/}identifier")
        comment_nodes = (
            ET.fromstring(z.read("word/comments.xml"))
            if "word/comments.xml" in z.namelist()
            else []
        )
        comments = {
            n.get(f"{{{W}}}id"): {
                "id": n.get(f"{{{W}}}id"),
                "author": n.get(f"{{{W}}}author"),
                "date": n.get(f"{{{W}}}date"),
                "text": text_of(n),
                "paragraphs": [],
            }
            for n in comment_nodes
        }
        paragraphs, active = [], set()
        for index, p in enumerate(root.findall(".//w:body//w:p", NS)):
            bookmarks = [
                b.get(f"{{{W}}}name", "") for b in p.findall(".//w:bookmarkStart", NS)
            ]
            bookmark = next((b for b in bookmarks if re.fullmatch(r"MMR\d+", b)), None)
            revisions = [
                {
                    "kind": r.tag.split("}")[-1],
                    "author": r.get(f"{{{W}}}author"),
                    "date": r.get(f"{{{W}}}date"),
                    "text": text_of(r, "all"),
                }
                for r in p.iter()
                if r.tag
                in (f"{{{W}}}ins", f"{{{W}}}del", f"{{{W}}}moveFrom", f"{{{W}}}moveTo")
            ]
            anchored = set(active)
            for el in p.iter():
                cid = el.get(f"{{{W}}}id")
                if el.tag == f"{{{W}}}commentRangeStart":
                    active.add(cid)
                    anchored.add(cid)
                elif el.tag == f"{{{W}}}commentRangeEnd":
                    active.discard(cid)
                elif el.tag == f"{{{W}}}commentReference":
                    anchored.add(cid)
            item = {
                "index": index,
                "bookmark": bookmark,
                "text": text_of(p),
                "original_text": text_of(p, "original"),
                "revisions": revisions,
            }
            paragraphs.append(item)
            for cid in anchored:
                if cid in comments:
                    comments[cid]["paragraphs"].append(
                        {"index": index, "bookmark": bookmark, "text": item["text"]}
                    )
        return {
            "export_id": identifier.text if identifier is not None else None,
            "paragraphs": paragraphs,
            "comments": list(comments.values()),
        }


def compare(args):
    reviewed = read_docx(args.reviewed)
    if args.baseline:
        baseline = args.baseline
    else:
        matches = list(
            (ROOT / "review-baselines").glob(f"*/{reviewed['export_id']}/manifest.json")
        )
        if len(matches) != 1:
            raise SystemExit(
                "Cannot identify one baseline; supply --baseline PATH/manifest.json."
            )
        baseline = matches[0]
    manifest = json.loads(baseline.read_text())
    if reviewed["export_id"] != manifest["export_id"]:
        raise SystemExit(
            "Export identifier mismatch: refusing to mix manuscript versions."
        )
    baseline_docx = baseline.parent / "baseline.docx"
    old = read_docx(baseline_docx)["paragraphs"]
    source_map = {p["bookmark"]: p for p in manifest["paragraphs"] if p["bookmark"]}
    changes, conflicts = [], []
    file_state = {}
    for path, snapshot in manifest["snapshots"].items():
        current = ROOT / path
        if not current.is_file() or digest(current.read_bytes()) != snapshot["sha256"]:
            file_state[path] = "source_changed_since_export"
    old_text, new_text = (
        [p["text"] for p in old],
        [p["text"] for p in reviewed["paragraphs"]],
    )
    matcher = difflib.SequenceMatcher(a=old_text, b=new_text, autojunk=False)
    for kind, a, b, c, d in matcher.get_opcodes():
        if kind == "equal":
            continue
        originals, replacements = old[a:b], reviewed["paragraphs"][c:d]
        hints = [
            source_map[p["bookmark"]] for p in originals if p["bookmark"] in source_map
        ]
        if not hints:
            hints = [
                source_map[p["bookmark"]]
                for p in old[max(0, a - 1) : a + 1]
                if p["bookmark"] in source_map
            ]
        paths = sorted({p["source"] for p in hints if p.get("source")})
        item = {
            "kind": kind,
            "before": originals,
            "after": replacements,
            "source_candidates": paths,
            "sections": sorted({p["section"] for p in hints}),
            "conflict": any(p in file_state for p in paths),
        }
        changes.append(item)
        if item["conflict"]:
            conflicts.append(item)
    for comment in reviewed["comments"]:
        comment["source_candidates"] = sorted(
            {
                source_map[p["bookmark"]]["source"]
                for p in comment["paragraphs"]
                if p["bookmark"] in source_map
                and source_map[p["bookmark"]].get("source")
            }
        )
    report = {
        "export_id": manifest["export_id"],
        "source_commit": manifest["source_commit"],
        "source_branch": manifest["source_branch"],
        "reviewed": str(args.reviewed.resolve()),
        "changes": changes,
        "comments": reviewed["comments"],
        "changed_source_files": file_state,
        "conflicts": len(conflicts),
        "policy": "Agent-assisted transfer only; no LaTeX files modified.",
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2))
    lines = [
        "# Word review transfer report",
        "",
        f"Source commit: `{manifest['source_commit']}`",
        f"Export: `{manifest['export_id']}`",
        "",
        f"Changed blocks: {len(changes)}; comments: {len(reviewed['comments'])}; "
        f"conflicts: {len(conflicts)}.",
        "",
        "Use the saved source snapshots for a three-way merge. Preserve original "
        "citation keys, labels, equations and generated inputs.",
        "",
    ]
    for i, change in enumerate(changes, 1):
        lines += [
            f"## Change {i}" + (" — SOURCE CONFLICT" if change["conflict"] else ""),
            "",
            "Source: " + ", ".join(change["source_candidates"]),
            "",
            "Before:",
            "",
            "\n\n".join(p["text"] for p in change["before"]),
            "",
            "After:",
            "",
            "\n\n".join(p["text"] for p in change["after"]),
            "",
        ]
    for comment in reviewed["comments"]:
        lines += [
            f"## Comment {comment['id']} by {comment['author']}",
            "",
            comment["text"],
            "",
            "Source: " + ", ".join(comment["source_candidates"]),
            "",
            "Anchor: " + " / ".join(p["text"] for p in comment["paragraphs"]),
            "",
        ]
    args.output.with_suffix(".md").write_text("\n".join(lines))
    print(
        json.dumps(
            {
                "report": str(args.output),
                "changed_blocks": len(changes),
                "comments": len(reviewed["comments"]),
                "conflicts": len(conflicts),
            },
            indent=2,
        )
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    exp = sub.add_parser("export")
    exp.add_argument("--manuscript", choices=("historical", "scenario"), required=True)
    exp.add_argument("--output", type=Path)
    exp.set_defaults(func=export)
    comp = sub.add_parser("compare")
    comp.add_argument("reviewed", type=Path)
    comp.add_argument("--baseline", type=Path)
    comp.add_argument(
        "--output", type=Path, default=ROOT / "build" / "word-review" / "report.json"
    )
    comp.set_defaults(func=compare)
    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
